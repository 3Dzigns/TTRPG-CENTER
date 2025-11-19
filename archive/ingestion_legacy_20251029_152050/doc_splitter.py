#!/usr/bin/env python3
"""
doc_splitter.py - Universal Document Splitter
==============================================

Splits documents by page range into a new output file.
Optionally updates Gate 0 marker files with split information.

Supported Formats:
  - PDF:  Physical pages (requires pypdf)
  - DOCX: Paragraph-based pages (50 paragraphs = 1 page)
  - TXT:  Line-based pages (50 lines = 1 page)

Usage:
  doc_splitter <source> <start_page> <end_page> <destination> [options]
  doc_splitter -v | --version
  doc_splitter -? | --help

Arguments:
  source       Path to source document
  start_page   Starting page number (1-based, inclusive)
  end_page     Ending page number (1-based, inclusive)
  destination  Path and filename for output document

Options:
  --update-marker FILE    Update Gate 0 marker file with split info
  --split-type TYPE       Split type: toc or part (required with --update-marker)
  --part-number N         Part number (required when --split-type is part)

Examples:
  # Basic split
  doc_splitter report.pdf 1 10 chapter1.pdf

  # Split TOC and update marker
  doc_splitter manual.pdf 1 8 manual_toc.pdf \
    --update-marker /Transfer_Station/Gate_0_Out/manual_20251009_120000.json \
    --split-type toc

  # Split Part 1 and update marker
  doc_splitter manual.pdf 9 100 manual_part1.pdf \
    --update-marker /Transfer_Station/Gate_0_Out/manual_20251009_120000.json \
    --split-type part --part-number 1

Version: 2.0.0
Author: n8n TTRPG Center
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Tuple, Dict, Any, Optional
from abc import ABC, abstractmethod


__version__ = "2.0.0"


# Configuration
LINES_PER_PAGE = 50      # For TXT files
PARAGRAPHS_PER_PAGE = 50 # For DOCX files


class DocumentSplitterError(Exception):
    """Base exception for document splitter errors."""
    pass


class BaseSplitter(ABC):
    """Abstract base class for document splitters."""

    def __init__(self, source_path: Path):
        self.source_path = source_path
        if not source_path.exists():
            raise DocumentSplitterError(f"Source file not found: {source_path}")

    @abstractmethod
    def get_total_pages(self) -> int:
        """Return total number of pages in document."""
        pass

    @abstractmethod
    def split(self, start_page: int, end_page: int, destination_path: Path) -> None:
        """Extract pages and write to destination."""
        pass

    def validate_range(self, start_page: int, end_page: int) -> Tuple[int, int]:
        """Validate and clamp page range to document bounds."""
        total_pages = self.get_total_pages()

        if start_page < 1:
            start_page = 1
        if end_page > total_pages:
            end_page = total_pages
        if start_page > end_page:
            raise DocumentSplitterError(
                f"Invalid range: start_page ({start_page}) > end_page ({end_page})"
            )
        if start_page > total_pages:
            raise DocumentSplitterError(
                f"Start page ({start_page}) exceeds total pages ({total_pages})"
            )

        return start_page, end_page


class PDFSplitter(BaseSplitter):
    """Split PDF documents by physical page numbers."""

    def __init__(self, source_path: Path):
        super().__init__(source_path)
        try:
            from pypdf import PdfReader, PdfWriter
            self.PdfReader = PdfReader
            self.PdfWriter = PdfWriter
        except ImportError:
            raise DocumentSplitterError(
                "pypdf library not installed. Run: pip install pypdf"
            )

        self.reader = self.PdfReader(str(self.source_path))

    def get_total_pages(self) -> int:
        return len(self.reader.pages)

    def split(self, start_page: int, end_page: int, destination_path: Path) -> None:
        start_page, end_page = self.validate_range(start_page, end_page)

        writer = self.PdfWriter()

        # Convert to 0-based indices
        for page_num in range(start_page - 1, end_page):
            writer.add_page(self.reader.pages[page_num])

        with open(destination_path, 'wb') as output_file:
            writer.write(output_file)


class DocxSplitter(BaseSplitter):
    """Split DOCX documents by paragraph-based pages."""

    def __init__(self, source_path: Path):
        super().__init__(source_path)
        try:
            from docx import Document
            self.Document = Document
        except ImportError:
            raise DocumentSplitterError(
                "python-docx library not installed. Run: pip install python-docx"
            )

        self.doc = self.Document(str(self.source_path))
        self.paragraphs = [p for p in self.doc.paragraphs]

    def get_total_pages(self) -> int:
        """Calculate pages based on paragraph count."""
        total_paragraphs = len(self.paragraphs)
        return (total_paragraphs + PARAGRAPHS_PER_PAGE - 1) // PARAGRAPHS_PER_PAGE

    def split(self, start_page: int, end_page: int, destination_path: Path) -> None:
        start_page, end_page = self.validate_range(start_page, end_page)

        # Convert pages to paragraph indices (0-based)
        start_para = (start_page - 1) * PARAGRAPHS_PER_PAGE
        end_para = end_page * PARAGRAPHS_PER_PAGE

        # Clamp to actual paragraph count
        total_paras = len(self.paragraphs)
        end_para = min(end_para, total_paras)

        # Create new document with selected paragraphs
        new_doc = self.Document()

        # Copy paragraphs
        for para in self.paragraphs[start_para:end_para]:
            new_para = new_doc.add_paragraph(para.text)
            # Preserve paragraph style
            if para.style:
                new_para.style = para.style.name

        new_doc.save(str(destination_path))


class TxtSplitter(BaseSplitter):
    """Split text documents by line-based pages."""

    def __init__(self, source_path: Path):
        super().__init__(source_path)
        with open(self.source_path, 'r', encoding='utf-8', errors='replace') as f:
            self.lines = f.readlines()

    def get_total_pages(self) -> int:
        """Calculate pages based on line count."""
        total_lines = len(self.lines)
        return (total_lines + LINES_PER_PAGE - 1) // LINES_PER_PAGE

    def split(self, start_page: int, end_page: int, destination_path: Path) -> None:
        start_page, end_page = self.validate_range(start_page, end_page)

        # Convert pages to line indices (0-based)
        start_line = (start_page - 1) * LINES_PER_PAGE
        end_line = end_page * LINES_PER_PAGE

        # Clamp to actual line count
        total_lines = len(self.lines)
        end_line = min(end_line, total_lines)

        # Write selected lines
        with open(destination_path, 'w', encoding='utf-8') as f:
            f.writelines(self.lines[start_line:end_line])


def get_splitter(source_path: Path) -> BaseSplitter:
    """Factory function to create appropriate splitter based on file extension."""
    extension = source_path.suffix.lower()

    splitter_map = {
        '.pdf': PDFSplitter,
        '.docx': DocxSplitter,
        '.txt': TxtSplitter,
    }

    splitter_class = splitter_map.get(extension)
    if not splitter_class:
        supported = ', '.join(splitter_map.keys())
        raise DocumentSplitterError(
            f"Unsupported file type: {extension}\n"
            f"Supported formats: {supported}"
        )

    return splitter_class(source_path)


def validate_marker_file(marker_path: Path) -> Dict[str, Any]:
    """
    Validate that marker file was created by gate_0_hash.py.

    Args:
        marker_path: Path to marker file

    Returns:
        Parsed marker data

    Raises:
        DocumentSplitterError: If marker file is invalid
    """
    if not marker_path.exists():
        raise DocumentSplitterError(f"Marker file not found: {marker_path}")

    try:
        with open(marker_path, 'r', encoding='utf-8') as f:
            marker_data = json.load(f)
    except json.JSONDecodeError as e:
        raise DocumentSplitterError(f"Invalid JSON in marker file: {e}")
    except IOError as e:
        raise DocumentSplitterError(f"Failed to read marker file: {e}")

    # Validate required fields (from gate_0_hash.py)
    required_fields = ['document_id', 'original_filename', 'sha256_hash']
    missing_fields = [field for field in required_fields if field not in marker_data]

    if missing_fields:
        raise DocumentSplitterError(
            f"Marker file missing required fields: {', '.join(missing_fields)}\n"
            f"This file was not created by gate_0_hash.py"
        )

    return marker_data


def update_marker_file(
    marker_path: Path,
    split_info: Dict[str, Any]
) -> None:
    """
    Update marker file with split information.

    Args:
        marker_path: Path to marker file
        split_info: Dict with split metadata

    Raises:
        DocumentSplitterError: If update fails
    """
    # Read current marker data
    marker_data = validate_marker_file(marker_path)

    # Initialize splits array if not present
    if 'splits' not in marker_data:
        marker_data['splits'] = []

    # Add new split info
    marker_data['splits'].append(split_info)

    # Write updated marker
    try:
        with open(marker_path, 'w', encoding='utf-8') as f:
            json.dump(marker_data, f, indent=2, ensure_ascii=False)
    except IOError as e:
        raise DocumentSplitterError(f"Failed to update marker file: {e}")


def parse_args() -> argparse.Namespace:
    """Parse command line arguments."""
    parser = argparse.ArgumentParser(
        description='Universal document splitter for PDF, DOCX, and TXT files',
        usage='%(prog)s <source> <start_page> <end_page> <destination> [--update-marker FILE] [--split-type {toc,part}] [--part-number N] [-v] [-?]',
        add_help=False  # We'll handle -? manually
    )

    # Positional arguments
    parser.add_argument(
        'source',
        nargs='?',
        type=Path,
        help='Source document path'
    )
    parser.add_argument(
        'start_page',
        nargs='?',
        type=int,
        help='Starting page number (1-based, inclusive)'
    )
    parser.add_argument(
        'end_page',
        nargs='?',
        type=int,
        help='Ending page number (1-based, inclusive)'
    )
    parser.add_argument(
        'destination',
        nargs='?',
        type=Path,
        help='Destination document path and filename'
    )

    # Marker update options
    parser.add_argument(
        '--update-marker',
        type=Path,
        metavar='FILE',
        help='Update Gate 0 marker file with split info'
    )
    parser.add_argument(
        '--split-type',
        type=str,
        choices=['toc', 'part'],
        help='Split type: toc or part (required with --update-marker)'
    )
    parser.add_argument(
        '--part-number',
        type=int,
        metavar='N',
        help='Part number (required when --split-type is part)'
    )

    # Version and help
    parser.add_argument(
        '-v', '--version',
        action='version',
        version=f'doc_splitter v{__version__}'
    )
    parser.add_argument(
        '-?', '--help',
        action='help',
        help='Show this help message and exit'
    )

    args = parser.parse_args()

    # Validate required arguments
    if not all([args.source, args.start_page, args.end_page, args.destination]):
        parser.print_help()
        sys.exit(1)

    # Validate marker update arguments
    if args.update_marker:
        if not args.split_type:
            parser.error("--split-type is required when using --update-marker")
        if args.split_type == 'part' and args.part_number is None:
            parser.error("--part-number is required when --split-type is part")

    return args


def main() -> int:
    """Main entry point."""
    try:
        args = parse_args()

        # Get appropriate splitter
        splitter = get_splitter(args.source)

        # Perform split operation
        splitter.split(args.start_page, args.end_page, args.destination)

        # Success message
        source_name = args.source.name
        dest_name = args.destination.name
        print(f"{source_name} pages {args.start_page} to {args.end_page} extracted to {dest_name}")

        # Update marker file if requested
        if args.update_marker:
            # Build split info
            split_info = {
                "type": args.split_type,
                "filename": dest_name,
                "path": str(args.destination.absolute()),
                "pages": f"{args.start_page}-{args.end_page}",
                "created_at": datetime.utcnow().isoformat() + 'Z'
            }

            # Add part_number for part splits
            if args.split_type == 'part':
                split_info["part_number"] = args.part_number

            # Update marker
            update_marker_file(args.update_marker, split_info)
            print(f"Updated marker: {args.update_marker.name}")

        return 0

    except DocumentSplitterError as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nOperation cancelled by user", file=sys.stderr)
        return 130
    except Exception as e:
        print(f"Unexpected error: {e}", file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
